from fastapi import UploadFile
from app.utils.deps import DB 
from app.services.imports.base_import_service import BaseImportService
from io import BytesIO
import pandas as pd
import math
from openpyxl import load_workbook

from app.models import (
    Product,
    Supplier, 
    Purchasing, 
    PurchasingDetail,
    Account
)

from app.utils.normalise import normalise_product_name, normalise_supplier_name, normalise_account_name
from app.utils.safe_parse import safe_str, safe_date, safe_number
from app.utils.cost_helper import update_avg_cost_for_products, refresh_product_avg_cost
from app.utils.event_flags import skip_cost_cache_updates
from app.utils.response import APIResponse

class LapPembelianImportService(BaseImportService):
    def __init__(self, db: DB):
        super().__init__(db)

    def _run(self, preview_id: str):
        payload = self.consume_preview(preview_id)
        rows = payload["rows"]
        add_products = payload["add_products"]
        update_products = payload["update_products"]
        suppliers = payload["suppliers"]
        accounts = payload["accounts"]

        purchasings_map = {}
        affected_product_ids = set()

        account_added = 0
        supplier_added = 0
        prod_added = 0
        prod_updated = 0
        purchasing_added = 0
        purchasing_skipped = []

        with skip_cost_cache_updates():
            for account in accounts:
                account_found = self.db.query(Account).filter_by(name=account["name"]).first()
                if not account_found:
                    account = Account(
                        name=account["name"]
                    )
                    self.db.add(account)
                    account_added += 1
                    self.db.flush()

            for supplier in suppliers:
                supplier_found = self.db.query(Supplier).filter_by(code=supplier["code"]).first()
                if not supplier_found:
                    supplier = Supplier(
                        code=supplier["code"],
                        name=supplier["name"],
                    )
                    self.db.add(supplier)
                    supplier_added += 1
                    self.db.flush()

            for prod in add_products:
                account_id = None

                if prod.get("account"):
                    account = (
                        self.db.query(Account)
                        .filter_by(name=prod["account"])
                        .first()
                    )
                    if account:
                        account_id = account.id

                product_found = self.db.query(Product).filter_by(name=prod["name"].upper()).first()
                if not product_found:
                    product = Product(
                        name=prod["name"].upper(),
                        unit=prod["unit"],
                        account_id=account_id
                    )
                    self.db.add(product)
                    prod_added += 1
                    self.db.flush()

            for prod in update_products:
                if not prod.get("account"):
                    continue

                account = (
                    self.db.query(Account)
                    .filter_by(name=prod["account"])
                    .first()
                )
                if not account:
                    continue  # should not happen, but stay safe

                db_product = (
                    self.db.query(Product)
                    .filter_by(id=prod["id"])
                    .first()
                )
                if not db_product:
                    continue

                if not db_product.account_id:
                    db_product.account_id = account.id
                    prod_updated += 1

            for row in rows:
                # --- collect required fields ---
                kode_supplier = row["supplier"] 
                tanggal = safe_date(row.get("tanggal"))
                no_bukti = row["no_bukti"]

                supplier = self.db.query(Supplier).filter_by(code=row["supplier"]).first()
                product = self.db.query(Product).filter_by(name=row["product"].upper()).first()

                if not supplier or not product:
                    purchasing_skipped.append(row)
                    # Should not happen, but never trust imports
                    continue

                # --- purchasing key ---
                key = (no_bukti if no_bukti else tanggal, kode_supplier)
                if key not in purchasings_map:
                    purchasing = Purchasing(
                        date=tanggal,
                        code=no_bukti,
                        purchase_order=row["purchase_order"],
                        supplier_id=supplier.id
                    )
                    self.db.add(purchasing)
                    self.db.flush()
                    purchasings_map[key] = purchasing
                else:
                    purchasing = purchasings_map[key]

                # --- detail insert ---
                detail = PurchasingDetail(
                    quantity=row["qty"],
                    price=row["price"],
                    discount=row["discount"],
                    ppn=row["ppn"],
                    dpp=row["dpp"],
                    pph=row["pph"],
                    tax_no=row["tax_no"],
                    exchange_rate=row["exchange_rate"] or 1,
                    product_id=product.id,
                    purchasing_id=purchasing.id
                )
                self.db.add(detail)
                affected_product_ids.add(product.id)
                purchasing_added += 1

            # self.db.commit()

        # Bulk update the cost cache for all affected products
        if affected_product_ids:
            update_avg_cost_for_products(self.db.connection(), list(affected_product_ids))

        # refresh_product_avg_cost(self.db)

        def sanitize(obj):
            if isinstance(obj, float) and (math.isnan(obj) or math.isinf(obj)):
                return None
            if isinstance(obj, dict):
                return {k: sanitize(v) for k, v in obj.items()}
            if isinstance(obj, list):
                return [sanitize(v) for v in obj]
            return obj
        
        return APIResponse.ok(data={
            "account_added": account_added,
            "supplier_added": supplier_added,
            "product_added": prod_added,
            "product_updated": prod_updated,
            "purchasing_added": purchasing_added,
            "purchasing_skipped": purchasing_skipped
        })

    def preview(self, file: UploadFile):
        contents: bytes = file.file.read()
        wb = load_workbook(BytesIO(contents), data_only=True, keep_links=False)

        HEADER_ROW = 6
        ROW_OFFSET = HEADER_ROW + 1

        valid_rows_data = []
        add_products = []
        update_products = []
        accounts = []
        suppliers = []

        summary = {"sheets": {}, "skipped": []}
        EXCLUDE_SHEETS = ["JANUARI 2025", "FEB 2025", "MARET", "APRIL", "MEI", "JUNI", "JULI"] # TODO: 2026 Change this shit
        for sheet in wb.sheetnames:
            if sheet in EXCLUDE_SHEETS:
                continue  

            df = pd.read_excel(BytesIO(contents), sheet_name=sheet, header=HEADER_ROW)
            df = df.iloc[:, :-2]

            rows_preview = []
            valid_rows = 0

            for idx, row in df.iterrows():
                excel_row_num = idx + ROW_OFFSET
                product_name_raw = row.get("NAMA BARANG")

                if pd.isna(product_name_raw):
                    continue

                product_name = safe_str(normalise_product_name(product_name_raw))
                if not product_name or product_name in {"NAT", "NONE"}:
                    continue

                kode_supplier = str(row.get("KODE SUPPLIER") or "").strip().upper()
                supplier_name = normalise_supplier_name(row.get("SUPPLIER") or "")
                tanggal = safe_date(row.get("TANGGAL"))
                no_bukti = safe_str(row.get("NO.BUKTI"))

                missing_cols = []
                if not kode_supplier:
                    missing_cols.append("KODE SUPPLIER")
                if not tanggal and not no_bukti:
                    missing_cols.append("TANGGAL/NO.BUKTI")

                if missing_cols:
                    summary["skipped"].append({
                        "sheet": sheet,
                        "row": excel_row_num,
                        "reason": f"missing column(s): {', '.join(missing_cols)}",
                        "product": product_name,
                    })
                    continue

                purchasing_found = self.db.query(Purchasing).filter_by(
                    code=no_bukti,
                    date=tanggal,
                ).first()
                if purchasing_found:
                    summary["skipped"].append({
                        "sheet": sheet,
                        "row": excel_row_num,
                        "reason": "duplicate NO.BUKTI",
                        "product": product_name,
                    })
                    continue

                supplier = self.db.query(Supplier).filter_by(code=kode_supplier).first()
                if not supplier:
                    suppliers.append({"code": kode_supplier, "name": supplier_name})

                acc_name_raw = row.get("ACCOUNT")
                acc_name_norm = normalise_account_name(acc_name_raw) if pd.notna(acc_name_raw) else None

                account_exists = None
                if acc_name_norm:
                    account_exists = (
                        self.db.query(Account)
                        .filter_by(name=acc_name_norm)
                        .first()
                    )

                if acc_name_norm:
                    already_in_preview = any(a["name"] == acc_name_norm for a in accounts)
                    if not account_exists and not already_in_preview:
                        accounts.append({"name": acc_name_norm})

                product = self.db.query(Product).filter_by(name=product_name.upper()).first()
                if not product:
                    add_products.append({
                        "name": product_name,
                        "unit": safe_str(row.get("SATUAN")),
                        "account": acc_name_norm
                    })
                elif not product.account_id:
                    update_products.append({
                        "id": product.id,
                        "account": acc_name_norm
                    })

                valid_rows += 1
                
                qty = safe_number(row.get("QTY")) or 0
                price = safe_number(row.get("HARGA SAT")) or 0
                ppn = safe_number(row.get("PPN")) or 0
                pph = safe_number(row.get("PPH")) or 0
                dpp = safe_number(row.get("DPP")) or 0
                
                res = {
                    "sheet": sheet,
                    "row": excel_row_num,
                    "supplier": kode_supplier,
                    "product": product_name,
                    "qty": qty,
                    "price": price,
                    "total": round(qty * price, 2),
                    "tanggal": tanggal.isoformat() if tanggal else None,
                    "no_bukti": no_bukti,
                    "purchase_order": safe_str(row.get("NO.PO")),
                    "discount": safe_number(row.get("POT.")) or 0.0,
                    "ppn": ppn / (qty or 1),
                    "dpp": dpp,
                    "pph": pph / (qty or 1),
                    "tax_no": safe_str(row.get("FAKTUR PAJAK")),
                    "exchange_rate": safe_number(row.get("KURS")) or 1,
                }

                rows_preview.append(res)

                valid_rows_data.append(res)

            summary["sheets"][sheet] = {
                "valid_rows": valid_rows,
                "preview_rows": rows_preview[:30],  # limit to first 30 rows per sheet
            }

        preview_id = self.create_preview({
            "rows": valid_rows_data,
            "add_products": add_products,
            "update_products": update_products,
            "accounts": accounts,
            "suppliers": suppliers
        })

        return APIResponse.ok(
            data={
                "preview_id": preview_id,
                "summary": {
                    "total_sheets": len(summary["sheets"]),
                    "skipped_total": len(summary["skipped"]),
                },
                "sheets": summary["sheets"],
                "skipped_sample": summary["skipped"][:50],
            }
        )